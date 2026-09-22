# -*- coding: utf-8 -*-
"""文档结构化抽取(services/extract.py)离线测试(不联网,不读写真实 data)。

    python -X utf8 tests/test_extract.py

覆盖 TODO F14 第 5 项(表头+样例、内嵌 JPEG 的 MIME、DOCX 取图失败走 note、
PDF 内嵌图与扫描件配额合计、status=ok 时图片 note 不被丢弃),对应本轮修复:
  - F5 PDF 内嵌图 + 扫描件各按 limit 计数(总计不超过 IMG_MAX_PER_FILE);
  - F6 内嵌图片按真实格式/文件名映射 MIME(不再硬编码 image/png);
  - F7 DOCX 同时覆盖 inline 与浮动(锚定)图片,按 rId 去重。
缺失可选依赖(Pillow / fitz / pypdf)的用例会打印 SKIP 并说明原因。
"""
import os
import shutil
import struct
import sys
import tempfile
import types
import zlib

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
if BACKEND not in sys.path:
    sys.path.insert(0, BACKEND)

from app.core import config as C              # noqa: E402
from app.services import extract as X          # noqa: E402

PASS, FAIL, SKIP = [], [], []


def _make_png_1x1(rgb=(255, 0, 0)):
    """手工构造合法 1x1 RGB PNG(不依赖 Pillow),供内嵌图用例使用。"""
    def _chunk(typ, data):
        return (struct.pack(">I", len(data)) + typ + data
                + struct.pack(">I", zlib.crc32(typ + data) & 0xffffffff))
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00" + bytes(rgb))         # 过滤字节 + RGB
    return (b"\x89PNG\r\n\x1a\n" + _chunk(b"IHDR", ihdr)
            + _chunk(b"IDAT", idat) + _chunk(b"IEND", b""))


def chk(tag, cond, detail=""):
    (PASS if cond else FAIL).append(tag)
    print("  %s %-46s %s" % ("[OK]" if cond else "[XX]", tag, detail))


def skip(tag, why):
    SKIP.append((tag, why))
    print("  [SKIP] %-44s %s" % (tag, why))


def _try_import(name):
    try:
        return __import__(name)
    except Exception:
        return None


class FakeVision(object):
    """假视觉 provider:记录每次调用的 (mime, label),可模拟不可用/识别失败。"""

    def __init__(self, available=True, reply="识别到:合同表单 字段A 状态:进行中",
                 raise_error=False):
        self.available = available
        self.reply = reply
        self.raise_error = raise_error
        self.calls = []                       # [(mime, label), ...]

    def complete_vision(self, sys_prompt, user_prompt, data_url):
        mime = ""
        if data_url.startswith("data:") and ";base64," in data_url:
            mime = data_url[5:].split(";", 1)[0]
        label = (user_prompt or "").split("\n", 1)[0]
        if label.startswith("文件名相关:"):
            label = label[len("文件名相关:"):]
        self.calls.append((mime, label))
        if self.raise_error:
            raise RuntimeError("mock vision failure")
        return self.reply

    @property
    def mimes(self):
        return [m for m, _ in self.calls]


def _docx_with_images(path, n=2, text="正文说明"):
    """生成含 n 张**互不相同**内嵌图的 docx(python-docx 会按内容去重同图)。"""
    from docx import Document
    doc = Document()
    if text:
        doc.add_paragraph(text)
    colors = [(255, 0, 0), (0, 255, 0), (0, 0, 255), (255, 255, 0),
              (0, 255, 255), (255, 0, 255), (128, 128, 0), (0, 128, 128)]
    for i in range(n):
        img = "%s_%d.png" % (path, i)
        with open(img, "wb") as f:
            f.write(_make_png_1x1(colors[i % len(colors)]))
        doc.add_picture(img)
    doc.save(path)


def main():
    tmp = tempfile.mkdtemp(prefix="h3ac_extract_test_")
    PIL = _try_import("PIL.Image")
    fitz = _try_import("fitz")
    pypdf = _try_import("pypdf")
    try:
        # ========== 0. 版本常量(缓存失效契约) ==========
        chk("VERSION 抽取逻辑版本为 2", getattr(X, "VERSION", None) == 2,
            str(getattr(X, "VERSION", None)))

        # ========== 1. xlsx 表头 + 样例(F14-5 第 1 项) ==========
        from openpyxl import Workbook
        p1 = os.path.join(tmp, "t1.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "合同"
        ws.append(["项目名称", "金额", "状态"])
        ws.append(["A项目", "100", "进行中"])
        ws.append(["B项目", "200", "完成"])
        wb.save(p1)
        wb.close()
        ext1 = X.extract_document(
            {"filename": "t1.xlsx", "stored_path": p1, "ext": ".xlsx"},
            provider=FakeVision(available=False))
        t = (ext1.get("tables") or [{}])[0]
        chk("F14-1 xlsx status=ok", ext1["status"] == "ok", ext1["status"])
        chk("F14-1 xlsx 列名(字段清单)",
            t.get("columns") == ["项目名称", "金额", "状态"], str(t.get("columns")))
        chk("F14-1 xlsx 样例数据行",
            t.get("rows") == [["A项目", "100", "进行中"], ["B项目", "200", "完成"]],
            str(t.get("rows")))
        chk("F14-1 xlsx rowCount 正确", t.get("rowCount") == 2, str(t.get("rowCount")))

        # ========== 2. 内嵌图片 MIME(F6) ==========
        # 2a. 辅助函数单元断言(不依赖 Pillow)
        chk("F6 _mime_for_format('jpeg')=image/jpeg",
            X._mime_for_format("jpeg") == "image/jpeg", X._mime_for_format("jpeg"))
        chk("F6 _mime_for_format('PNG') 忽略大小写",
            X._mime_for_format("PNG") == "image/png", X._mime_for_format("PNG"))
        chk("F6 _mime_for_format 未知格式兜底 image/png",
            X._mime_for_format("bogus") == "image/png", X._mime_for_format("bogus"))
        chk("F6 _mime_for_name('Im0.jpeg')=image/jpeg",
            X._mime_for_name("Im0.jpeg") == "image/jpeg", X._mime_for_name("Im0.jpeg"))
        chk("F6 _mime_for_name 大写扩展名 .JPG",
            X._mime_for_name("Im0.JPG") == "image/jpeg", X._mime_for_name("Im0.JPG"))
        chk("F6 _mime_for_name 无扩展名兜底 image/png",
            X._mime_for_name("Im0") == "image/png", X._mime_for_name("Im0"))

        # 2b. xlsx 内嵌 JPEG → 视觉调用收到 image/jpeg(需 Pillow 造真实 JPEG)
        if PIL is None:
            skip("F6 xlsx 内嵌 JPEG 实测", "缺少 Pillow,无法生成真实 JPEG")
        else:
            from PIL import Image as PILImage
            from openpyxl.drawing.image import Image as XLImage
            jpg = os.path.join(tmp, "pic.jpg")
            PILImage.new("RGB", (32, 32), (200, 30, 30)).save(jpg, "JPEG")
            with open(jpg, "rb") as f:
                magic = f.read(4).hex()
            p2 = os.path.join(tmp, "t2.xlsx")
            wb2 = Workbook()
            ws2 = wb2.active
            ws2["A1"] = "占位"
            ws2.add_image(XLImage(jpg), "D2")
            wb2.save(p2)
            wb2.close()
            prov2 = FakeVision(available=True)
            ext2 = X.extract_document(
                {"filename": "t2.xlsx", "stored_path": p2, "ext": ".xlsx"}, provider=prov2)
            chk("F6 xlsx 内嵌图为 JPEG(魔数 ffd8ffe0)", magic == "ffd8ffe0", magic)
            chk("F6 xlsx 视觉调用收到 image/jpeg",
                prov2.mimes == ["image/jpeg"], str(prov2.mimes))
            chk("F6 xlsx 图片识别结果进入 images.text",
                any(i.get("text") for i in ext2["images"]), str(ext2["images"]))

        # ========== 3. DOCX 图片/浮动图片 + 失败走 note(F7) ==========
        docx_path = os.path.join(tmp, "t3.docx")
        _docx_with_images(docx_path, n=2, text="这是正文")
        # 3a. _docx_pics 遍历 a:blip 且按 rId 去重(内联 2 张只应算 2,
        #     并额外验证"仅存在于 body、不在 inline_shapes"的浮动 blip 也能被收上来)
        from docx import Document as DocxDocument
        from lxml import etree
        W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        A = "http://schemas.openxmlformats.org/drawingml/2006/main"
        R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        body = etree.fromstring(
            ('<w:body xmlns:w="%s" xmlns:a="%s" xmlns:r="%s">'
             '<a:blip r:embed="rId1"/><a:blip r:embed="rId1"/>'
             '<a:blip r:embed="rId2"/></w:body>' % (W, A, R)).encode("utf-8"))

        class _FP(object):
            def __init__(self, blob, ct):
                self.blob, self.content_type = blob, ct

        fake_doc = types.SimpleNamespace(
            element=types.SimpleNamespace(body=body),
            part=types.SimpleNamespace(related_parts={
                "rId1": _FP(b"a", "image/png"),
                "rId2": _FP(b"b", "image/jpeg")}))
        pics = X._docx_pics(fake_doc)
        chk("F7 _docx_pics 遍历 a:blip(含浮动)且按 rId 去重",
            [rid for rid, _ in pics] == ["rId1", "rId2"], str([rid for rid, _ in pics]))
        real_pics = X._docx_pics(DocxDocument(docx_path))
        chk("F7 真实 docx 两张内联图只计数 2(不重复)",
            len(real_pics) == 2, "len=%d" % len(real_pics))
        # 3b. 未配置视觉模型 → note 且含图片数量
        ext3 = X.extract_document(
            {"filename": "t3.docx", "stored_path": docx_path, "ext": ".docx"},
            provider=FakeVision(available=False))
        notes = [i.get("note", "") for i in ext3["images"] if i.get("note")]
        chk("F7 未配置视觉:DOCX note 存在且含图片数量",
            bool(notes) and "2" in notes[0], str(notes))
        chk("F7 有正文时 status=ok(不被图片 note 拖成 empty)",
            ext3["status"] == "ok", ext3["status"])
        # 3c. 视觉识别失败 → 该图片条目走 note
        prov3 = FakeVision(available=True, raise_error=True)
        ext3b = X.extract_document(
            {"filename": "t3.docx", "stored_path": docx_path, "ext": ".docx"}, provider=prov3)
        chk("F7 视觉识别失败:DOCX 图片走 note",
            any(i.get("note") for i in ext3b["images"]), str(ext3b["images"]))

        # ========== 4. PDF 内嵌图 + 扫描件合计不超 limit(F5) ==========
        if PIL is None or fitz is None or pypdf is None:
            missing = [n for n, m in (("Pillow", PIL), ("fitz", fitz), ("pypdf", pypdf))
                       if m is None]
            skip("F5 PDF 配额实测", "缺少依赖:%s" % ",".join(missing))
        else:
            from PIL import Image as PILImage
            jpg = os.path.join(tmp, "pdf_pic.jpg")
            PILImage.new("RGB", (48, 48), (10, 120, 60)).save(jpg, "JPEG")
            pdf_path = os.path.join(tmp, "t4.pdf")
            d = fitz.open()
            for _ in range(2):                      # 2 页,每页 1 张内嵌图,无文字层
                pg = d.new_page()
                pg.insert_image(fitz.Rect(40, 40, 160, 160), filename=jpg)
            d.save(pdf_path)
            d.close()

            old_limit = C.IMG_MAX_PER_FILE
            # 4a. limit=2:两页内嵌图恰好占满 → 扫描件 remaining=0 必须跳过
            C.IMG_MAX_PER_FILE = 2
            try:
                prov4 = FakeVision(available=True)
                ext4 = X.extract_document(
                    {"filename": "t4.pdf", "stored_path": pdf_path, "ext": ".pdf"}, provider=prov4)
                chk("F5 limit=2:视觉调用=2(内嵌图占满,扫描件跳过)",
                    len(prov4.calls) == 2, "calls=%d" % len(prov4.calls))
                chk("F5 limit=2:len(images)<=limit",
                    len(ext4["images"]) <= 2, "images=%d" % len(ext4["images"]))
            finally:
                C.IMG_MAX_PER_FILE = old_limit
            # 4b. limit=3:内嵌 2 + 扫描件扣减后补 1
            C.IMG_MAX_PER_FILE = 3
            try:
                prov4b = FakeVision(available=True)
                ext4b = X.extract_document(
                    {"filename": "t4.pdf", "stored_path": pdf_path, "ext": ".pdf"}, provider=prov4b)
                chk("F5 limit=3:视觉调用=3(内嵌2 + 扫描件扣减补1)",
                    len(prov4b.calls) == 3, "calls=%d" % len(prov4b.calls))
                chk("F5 limit=3:len(images)<=limit",
                    len(ext4b["images"]) <= 3, "images=%d" % len(ext4b["images"]))
            finally:
                C.IMG_MAX_PER_FILE = old_limit

        # ========== 5. status=ok 时图片 note 仍保留(F14-5 第 5 项) ==========
        # 5a. _merge_notes:有正文时不产生整体 note,无正文时汇总图片 note
        chk("F14-5 有正文时 _merge_notes 不覆盖为整体错误 note",
            X._merge_notes(None, [], "有正文",
                           [{"from": "x", "note": "未识别图片"}], "docx") == "")
        agg = X._merge_notes(None, [], "", [{"from": "x", "note": "未识别图片"}], "docx")
        chk("F14-5 无正文时 note 汇总图片说明",
            "未识别图片" in agg, agg)
        # 5b. 真实 xlsx:有表格 + 未识别图片 → status=ok 且 images 里 note 未被丢弃
        if PIL is None:
            skip("F14-5 status=ok + 图片 note 实测", "缺少 Pillow,无法插入内嵌图")
        else:
            from PIL import Image as PILImage
            from openpyxl.drawing.image import Image as XLImage
            jpg = os.path.join(tmp, "note_pic.jpg")
            PILImage.new("RGB", (16, 16), (0, 0, 200)).save(jpg, "JPEG")
            p5 = os.path.join(tmp, "t5.xlsx")
            wb5 = Workbook()
            ws5 = wb5.active
            ws5.title = "数据"
            ws5.append(["字段", "类型"])
            ws5.append(["金额", "数字"])
            ws5.add_image(XLImage(jpg), "D2")
            wb5.save(p5)
            wb5.close()
            ext5 = X.extract_document(
                {"filename": "t5.xlsx", "stored_path": p5, "ext": ".xlsx"},
                provider=FakeVision(available=False))
            chk("F14-5 有表格+未识别图片时 status=ok",
                ext5["status"] == "ok", ext5["status"])
            chk("F14-5 图片 note 未被丢弃(仍在 images 中)",
                any(i.get("note") for i in ext5["images"]), str(ext5["images"]))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print("\n" + "=" * 66)
    print("文档抽取测试:通过 %d / 失败 %d / 跳过 %d" % (len(PASS), len(FAIL), len(SKIP)))
    if FAIL:
        print("失败项:", FAIL)
    for tag, why in SKIP:
        print("  SKIP %s → %s" % (tag, why))
    print("=" * 66)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())

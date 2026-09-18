# -*- coding: utf-8 -*-
"""文档解析:parse_file(path) -> str。

按扩展名分流:
  .xlsx/.xlsm  openpyxl(含所有 sheet,行用 " | " 连接)
  .docx        python-docx(段落 + 表格)
  .pdf         pypdf(逐页 extract_text)
  其余文本     多编码尝试(utf-8/utf-8-sig/gbk/gb18030/latin-1)
异常统一返回 "";结果截断 200000 字符并压缩连续空行。
"""
import os
import re

MAX_CHARS = 200_000
_TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gbk", "gb18030", "latin-1")


def _clean(text):
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n").lstrip("\ufeff")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_CHARS]


def _parse_xlsx(path):
    from openpyxl import load_workbook
    parts = []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            lines = ["# 工作表:%s" % ws.title]
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    lines.append(" | ".join(cells))
            parts.append("\n".join(lines))
    finally:
        wb.close()
    return "\n\n".join(parts)


def _parse_docx(path):
    from docx import Document
    doc = Document(path)
    parts = [(p.text or "").strip() for p in doc.paragraphs if (p.text or "").strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _parse_text(path):
    with open(path, "rb") as f:
        raw = f.read()
    for enc in _TEXT_ENCODINGS:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", "replace")


def parse_file(path: str) -> str:
    if not path or not os.path.isfile(path):
        return ""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            text = _parse_xlsx(path)
        elif ext == ".docx":
            text = _parse_docx(path)
        elif ext == ".pdf":
            text = _parse_pdf(path)
        else:
            text = _parse_text(path)
    except Exception:
        return ""
    return _clean(text)
